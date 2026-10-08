"""Excel and CSV export.

The workbook keeps inputs as values and writes subtotals, net cash flow and the
running cash balance as live formulas, so the statements stay editable in Excel.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .engine import CATEGORIES, SECTIONS
from .models import Assumptions

MONEY = '#,##0;[Red](#,##0);"-"'
PCT = "0.0%"
DATE = "yyyy-mm-dd"
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
SECTION_FILL = PatternFill("solid", fgColor="E5E7EB")
TOTAL_FILL = PatternFill("solid", fgColor="DBEAFE")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="9CA3AF")


def _header(ws, row: int, values: List[Any], start_col: int = 1) -> None:
    for i, value in enumerate(values):
        cell = ws.cell(row=row, column=start_col + i, value=value)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center" if i else "left")


def _widths(ws, first: float, rest: float, columns: int) -> None:
    ws.column_dimensions["A"].width = first
    for col in range(2, columns + 1):
        ws.column_dimensions[get_column_letter(col)].width = rest


def _statement(ws, title: str, periods: List[Dict[str, Any]], starting_cash: float, min_cash: float) -> None:
    ncols = len(periods)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    _header(ws, 3, ["Cash flow (USD)"] + [p["label"] for p in periods])
    ws.cell(row=4, column=1, value="Period start").font = Font(italic=True)
    ws.cell(row=5, column=1, value="Period end (exclusive)").font = Font(italic=True)
    for i, p in enumerate(periods):
        for r, key in ((4, "start"), (5, "end")):
            cell = ws.cell(row=r, column=2 + i, value=p[key])
            cell.number_format = DATE
            cell.font = Font(italic=True)
            cell.alignment = Alignment(horizontal="center")

    row = 7
    section_rows: Dict[str, int] = {}
    for section in SECTIONS:
        ws.cell(row=row, column=1, value=section.upper() + " ACTIVITIES").font = BOLD
        for col in range(1, ncols + 2):
            ws.cell(row=row, column=col).fill = SECTION_FILL
        row += 1
        first = row
        for key, label, sec in CATEGORIES:
            if sec != section:
                continue
            ws.cell(row=row, column=1, value=label)
            for i, p in enumerate(periods):
                cell = ws.cell(row=row, column=2 + i, value=round(p["categories"][key], 2))
                cell.number_format = MONEY
            row += 1
        last = row - 1
        ws.cell(row=row, column=1, value=f"Net {section} cash flow").font = BOLD
        for i in range(ncols):
            col = get_column_letter(2 + i)
            cell = ws.cell(row=row, column=2 + i, value=f"=SUM({col}{first}:{col}{last})")
            cell.number_format = MONEY
            cell.font = BOLD
            cell.border = Border(top=THIN)
        section_rows[section] = row
        row += 2

    net_row = row
    ws.cell(row=net_row, column=1, value="NET CASH FLOW").font = BOLD
    begin_row, end_row, min_row, flag_row = net_row + 2, net_row + 3, net_row + 4, net_row + 5
    ws.cell(row=begin_row, column=1, value="Beginning cash")
    ws.cell(row=end_row, column=1, value="ENDING CASH").font = BOLD
    ws.cell(row=min_row, column=1, value="Minimum cash target")
    ws.cell(row=flag_row, column=1, value="Below minimum?")
    for i in range(ncols):
        col = get_column_letter(2 + i)
        prev = get_column_letter(1 + i)
        net = ws.cell(
            row=net_row,
            column=2 + i,
            value="=" + "+".join(f"{col}{section_rows[s]}" for s in SECTIONS),
        )
        net.number_format = MONEY
        net.font = BOLD
        net.fill = TOTAL_FILL
        begin = ws.cell(
            row=begin_row,
            column=2 + i,
            value=round(starting_cash, 2) if i == 0 else f"={prev}{end_row}",
        )
        begin.number_format = MONEY
        end = ws.cell(row=end_row, column=2 + i, value=f"={col}{begin_row}+{col}{net_row}")
        end.number_format = MONEY
        end.font = BOLD
        end.fill = TOTAL_FILL
        end.border = Border(top=THIN, bottom=THIN)
        minimum = ws.cell(row=min_row, column=2 + i, value=round(min_cash, 2))
        minimum.number_format = MONEY
        ws.cell(row=flag_row, column=2 + i, value=f'=IF({col}{end_row}<{col}{min_row},"YES","")').alignment = Alignment(horizontal="center")
    ws.cell(row=net_row, column=1).fill = TOTAL_FILL
    ws.cell(row=end_row, column=1).fill = TOTAL_FILL
    ws.freeze_panes = "B6"
    _widths(ws, 38, 14, ncols + 1)


def _pnl(ws, pnl: List[Dict[str, Any]]) -> None:
    ws["A1"] = "Monthly profit & loss (accrual basis, before income tax)"
    ws["A1"].font = Font(bold=True, size=14)
    _header(ws, 3, ["USD"] + [p["label"] for p in pnl])
    lines = [
        ("Revenue", "revenue"),
        ("Cost of sales", "cogs"),
        ("Gross profit", None),
        ("Payroll & benefits", "payroll"),
        ("Operating expenses", "opex"),
        ("Interest expense", "interest"),
        ("Depreciation (non-cash)", "depreciation"),
        ("Pre-tax profit", None),
    ]
    for r, (label, key) in enumerate(lines, start=4):
        ws.cell(row=r, column=1, value=label)
        for i, p in enumerate(pnl):
            col = get_column_letter(2 + i)
            if key:
                value: Any = round(p[key], 2)
            elif label == "Gross profit":
                value = f"={col}4-{col}5"
            else:
                value = f"={col}6-{col}7-{col}8-{col}9-{col}10"
            cell = ws.cell(row=r, column=2 + i, value=value)
            cell.number_format = MONEY
            if not key:
                cell.font = BOLD
                cell.fill = TOTAL_FILL
        if not key:
            ws.cell(row=r, column=1).font = BOLD
            ws.cell(row=r, column=1).fill = TOTAL_FILL
    _widths(ws, 30, 13, len(pnl) + 1)
    ws.freeze_panes = "B4"


def _assumptions(ws, a: Assumptions) -> None:
    ws["A1"] = "Assumptions"
    ws["A1"].font = Font(bold=True, size=14)
    data = a.model_dump(mode="json")
    row = 3
    for section in ("general", "sales", "costs", "ap", "macro"):
        ws.cell(row=row, column=1, value=section.title()).font = BOLD
        row += 1
        for key, value in data[section].items():
            ws.cell(row=row, column=1, value=key.replace("_", " "))
            ws.cell(row=row, column=2, value=value)
            row += 1
        row += 1
    ws.cell(row=row, column=1, value="Payroll").font = BOLD
    row += 1
    for key, value in data["payroll"].items():
        if key in ("hires", "employees"):
            continue
        ws.cell(row=row, column=1, value=key.replace("_", " "))
        ws.cell(row=row, column=2, value=value)
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Receivable collectability (%) and overdue lag (days), by aging").font = BOLD
    row += 1
    _header(ws, row, ["Bucket", "Collectability %", "Lag days"])
    row += 1
    for bucket, pct in data["collections"]["collectability_pct"].items():
        ws.cell(row=row, column=1, value=bucket)
        ws.cell(row=row, column=2, value=pct)
        ws.cell(row=row, column=3, value=data["collections"]["overdue_lag_days"][bucket])
        row += 1
    row += 1

    def table(title: str, headers: List[str], items: List[List[Any]]) -> None:
        nonlocal row
        ws.cell(row=row, column=1, value=title).font = BOLD
        row += 1
        _header(ws, row, headers)
        row += 1
        for item in items:
            for col, value in enumerate(item, start=1):
                ws.cell(row=row, column=col, value=value)
            row += 1
        row += 1

    table(
        "Employee roster",
        ["ID", "Department", "Title", "Pay type", "Annual base", "Hire date", "Term date", "Bonus %", "Benefits / month"],
        [[e["id"], e["department"], e["title"], e["pay_type"],
          e["annual_salary"] if e["pay_type"] == "salary" else round(e["hourly_rate"] * e["hours_per_week"] * 52, 2),
          e["hire_date"], e["term_date"], e["bonus_pct"], e["benefits_monthly"]] for e in data["payroll"]["employees"]],
    )
    table("Hiring plan", ["Month #", "Heads"], [[h["month"], h["count"]] for h in data["payroll"]["hires"]])
    table(
        "Operating expense lines",
        ["Name", "Kind", "Amount", "Growth %/mo", "Start month", "End month"],
        [[o["name"], o["kind"], o["amount"], o["growth_pct_monthly"], o["start_month"], o["end_month"]] for o in data["opex"]],
    )
    table(
        "Loans",
        ["Name", "Balance", "Rate %", "Monthly payment"],
        [[l["name"], l["balance"], l["annual_rate_pct"], l["monthly_payment"]] for l in data["loans"]],
    )
    table(
        "One-time items",
        ["Name", "Date", "Amount", "Category"],
        [[o["name"], o["date"], o["amount"], o["category"]] for o in data["one_time"]],
    )
    _widths(ws, 34, 18, 6)


def _ar(ws, ar_data: Dict[str, Any]) -> None:
    ws["A1"] = "Receivables outlook"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"], ws["B3"] = "Open receivables", round(ar_data["open_total"], 2)
    ws["A4"], ws["B4"] = "Expected to collect (after haircuts)", round(ar_data["expected_total"], 2)
    ws["A5"], ws["B5"] = "Expected within the forecast horizon", round(ar_data["expected_in_horizon"], 2)
    for r in (3, 4, 5):
        ws.cell(row=r, column=2).number_format = MONEY

    _header(ws, 7, ["Aging bucket", "Open amount", "Invoices"])
    row = 8
    for item in ar_data["aging"]:
        ws.cell(row=row, column=1, value=item["label"])
        c = ws.cell(row=row, column=2, value=round(item["open_amount"], 2))
        c.number_format = MONEY
        ws.cell(row=row, column=3, value=item["invoices"])
        row += 1

    row += 1
    _header(
        ws, row,
        ["Customer", "Open invoices", "Open amount", "Expected in horizon", "Avg days to pay", "Oldest days past due", "Risk"],
    )
    row += 1
    for cust in ar_data["customers"]:
        values = [
            cust["customer_id"], cust["open_invoices"], round(cust["open_amount"], 2),
            round(cust["expected_in_horizon"], 2),
            None if cust["avg_days_to_pay"] is None else round(cust["avg_days_to_pay"], 1),
            cust["oldest_days_past_due"], cust["risk"],
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            if col in (3, 4):
                cell.number_format = MONEY
        row += 1
    _widths(ws, 36, 20, 7)


def _ap(ws, ap_data: Dict[str, Any]) -> None:
    ws["A1"] = "Payables outlook"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"], ws["B3"] = "Open vendor bills", round(ap_data["open_total"], 2)
    ws["A4"], ws["B4"] = "Of which past due", round(ap_data["overdue_total"], 2)
    ws["A5"], ws["B5"] = "Actual days payable (bill to payment)", ap_data["actual_dpo_days"]
    ws["B3"].number_format = ws["B4"].number_format = MONEY
    _header(ws, 7, ["Aging bucket", "Open amount", "Bills"])
    row = 8
    for item in ap_data["aging"]:
        ws.cell(row=row, column=1, value=item["label"])
        ws.cell(row=row, column=2, value=round(item["open_amount"], 2)).number_format = MONEY
        ws.cell(row=row, column=3, value=item["bills"])
        row += 1
    row += 1
    _header(ws, row, ["Vendor", "Category", "Open amount", "Share %", "Avg days to pay", "Avg days vs due", "Oldest days past due"])
    row += 1
    for v in ap_data["vendors"]:
        values = [v["vendor_name"], v["category"], round(v["open_amount"], 2), v["share_pct"],
                  v["avg_days_to_pay"], v["avg_days_vs_due"], v["oldest_days_past_due"]]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            if col == 3:
                cell.number_format = MONEY
        row += 1
    _widths(ws, 34, 20, 7)


def _capex(ws, capex_data: Dict[str, Any]) -> None:
    ws["A1"] = "Capital expenditure plan"
    ws["A1"].font = Font(bold=True, size=14)
    t = capex_data["totals"]
    for r, (label, key) in enumerate(
        [("Cash capex in horizon", "cash_capex"), ("  Maintenance", "maintenance"), ("  Growth", "growth"),
         ("Financed purchases", "financed_amount"), ("Financing payments in horizon", "financed_payments"),
         ("Depreciation in horizon (non-cash)", "depreciation")], start=3):
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=round(t[key], 2)).number_format = MONEY
    _header(ws, 10, ["Item", "Category", "Kind", "Funding", "Date", "Amount", "Cash at purchase", "Financed", "Monthly payment", "Depreciation / month"])
    row = 11
    for it in capex_data["items"]:
        values = [it["name"], it["category"], it["kind"], it["funding"], it["date"], round(it["amount"], 2),
                  round(it["cash_at_purchase"], 2), round(it["financed"], 2),
                  round(it["monthly_payment"], 2), round(it["depreciation_monthly"], 2)]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            if col >= 6:
                cell.number_format = MONEY
        row += 1
    row += 1
    _header(ws, row, ["Month", "Maintenance", "Growth", "Financing payments", "Interest", "Depreciation", "Net asset additions (cum.)"])
    row += 1
    for m in capex_data["monthly"]:
        values = [m["label"], m["maintenance"], m["growth"], m["financed_payments"], m["interest"], m["depreciation"], m["net_additions_cum"]]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=round(value, 2) if col > 1 else value)
            if col > 1:
                cell.number_format = MONEY
        row += 1
    _widths(ws, 30, 16, 10)


def _table_sheet(ws, title: str, headers: List[str], rows: List[List[Any]], first_width: float = 30, width: float = 15) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    _header(ws, 3, headers)
    for r, values in enumerate(rows, start=4):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=round(value, 2) if isinstance(value, float) else value)
            if isinstance(value, (int, float)) and c > 1:
                cell.number_format = MONEY
    _widths(ws, first_width, width, len(headers))


def _payroll(ws, pr: Dict[str, Any]) -> None:
    ws["A1"] = "Payroll"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"{pr['pay_frequency']} pay runs; cost is accrued by month, cash goes out on pay runs"
    labels = [f"M{i + 1}" for i in range(len(pr["monthly"]))]
    _header(ws, 4, ["Accrued cost"] + labels + ["Total"])
    r = 5
    for dept, vals in pr["by_department"].items():
        for c, v in enumerate([dept] + list(vals) + [sum(vals)], start=1):
            cell = ws.cell(row=r, column=c, value=round(v, 2) if c > 1 else v)
            if c > 1:
                cell.number_format = MONEY
        r += 1
    for label, vals in (("Total payroll cost", pr["monthly"]), ("  of which bonuses", pr["bonus"]), ("  of which benefits", pr["benefits"])):
        for c, v in enumerate([label] + list(vals) + [sum(vals)], start=1):
            cell = ws.cell(row=r, column=c, value=round(v, 2) if c > 1 else v)
            cell.font = BOLD if label.startswith("Total") else Font()
            if c > 1:
                cell.number_format = MONEY
        r += 1
    ws.cell(row=r, column=1, value="Headcount (end of month)")
    for c, v in enumerate(pr["headcount"], start=2):
        ws.cell(row=r, column=c, value=v)
    r += 2
    _header(ws, r, ["Pay run date", "Gross wages", "Employer tax", "Total cash", "Employees paid"])
    r += 1
    for run in pr["runs"]:
        for c, v in enumerate([run["date"], run["gross"], run["employer_tax"], run["total"], run["employees"]], start=1):
            cell = ws.cell(row=r, column=c, value=round(v, 2) if isinstance(v, float) else v)
            if c in (2, 3, 4):
                cell.number_format = MONEY
        r += 1
    _widths(ws, 30, 13, len(labels) + 2)


def _balance_sheet(ws, bs: Dict[str, Any]) -> None:
    cols = [bs["opening"]] + bs["months"]
    lines = [("Cash", "cash"), ("Receivables", "receivables"), ("Property & equipment (net)", "ppe_net"), ("Total assets", "total_assets"),
             ("Payables", "payables"), ("Accrued payroll", "accrued_payroll"), ("Taxes payable", "taxes_payable"), ("Debt", "debt"),
             ("Total liabilities", "total_liabilities"), ("Equity", "equity"), ("Total liabilities + equity", "total_liabilities_equity"),
             ("Check (assets - liabilities - equity)", "check")]
    _table_sheet(ws, "Projected balance sheet", ["USD"] + [c["label"] for c in cols],
                 [[label] + [c[key] for c in cols] for label, key in lines], 34, 13)
    ws.cell(row=len(lines) + 5, column=1, value="Receivables not expected to be collected (existing)")
    ws.cell(row=len(lines) + 5, column=2, value=round(bs["memo"]["existing_ar_expected_uncollectible"], 2)).number_format = MONEY


def _gl(ws, g: Dict[str, Any]) -> None:
    rows = []
    for t in g["timeline"]:
        rows.append([t["label"], t["actual_revenue"], t["actual_costs"], t["actual_pretax"],
                     t["forecast_revenue"], t["forecast_costs"], t["forecast_pretax"]])
    _table_sheet(ws, "Actuals (general ledger) vs forecast", ["Month", "Actual revenue", "Actual costs", "Actual pre-tax",
                                                              "Forecast revenue", "Forecast costs", "Forecast pre-tax"], rows, 18, 16)


def _projections(ws, result: Dict[str, Any]) -> None:
    ws["A1"] = "Standard projections: 30, 60, 90, 180 days and 1 year"
    ws["A1"].font = Font(bold=True, size=14)
    projections = result["projections"]
    headers = ["Measure"] + [p["label"] + ("" if p["complete"] else " (beyond horizon)") for p in projections]
    _header(ws, 3, headers)
    rows = [
        ("Ends on", [p["end_date"] for p in projections]),
        ("Starting cash", [p["starting_cash"] for p in projections]),
        ("Cash in", [p["cash_in"] for p in projections]),
        ("Cash out", [-p["cash_out"] for p in projections]),
        ("Operating cash flow", [p["operating"] for p in projections]),
        ("Investing cash flow", [p["investing"] for p in projections]),
        ("Financing cash flow", [p["financing"] for p in projections]),
        ("Net change in cash", [p["net_cash_flow"] for p in projections]),
        ("Ending cash", [p["ending_cash"] for p in projections]),
        ("Lowest balance", [p["lowest_balance"] for p in projections]),
        ("Lowest balance date", [p["lowest_balance_date"] for p in projections]),
        ("First below minimum", [p["first_below_min_date"] or "Never" for p in projections]),
        ("Funding gap vs minimum", [p["funding_gap"] for p in projections]),
    ]
    for r, (label, values) in enumerate(rows, start=4):
        ws.cell(row=r, column=1, value=label)
        for c, v in enumerate(values, start=2):
            if isinstance(v, (dt.date, dt.datetime)):
                v = v.isoformat()
            cell = ws.cell(row=r, column=c, value=round(v, 2) if isinstance(v, float) else v)
            if isinstance(v, (int, float)):
                cell.number_format = MONEY
    r = 4 + len(rows) + 1
    ws.cell(row=r, column=1, value="Ending cash by scenario").font = BOLD
    for i, (name, c) in enumerate(result["comparison"].items(), start=r + 1):
        ws.cell(row=i, column=1, value=c["label"])
        for j, v in enumerate(c["projection_end_cash"], start=2):
            ws.cell(row=i, column=j, value=round(v, 2)).number_format = MONEY
    _widths(ws, 30, 18, len(headers))


def _scenarios(ws, result: Dict[str, Any]) -> None:
    ws["A1"] = "Scenario comparison: ending cash by month"
    ws["A1"].font = Font(bold=True, size=14)
    labels = [p["label"] for p in result["monthly"]]
    _header(ws, 3, ["Scenario"] + labels)
    for r, (name, data) in enumerate(result["comparison"].items(), start=4):
        ws.cell(row=r, column=1, value=data["label"])
        for i, value in enumerate(data["monthly_end_cash"]):
            cell = ws.cell(row=r, column=2 + i, value=round(value, 2))
            cell.number_format = MONEY
    _widths(ws, 24, 14, len(labels) + 1)


def _summary(ws, result: Dict[str, Any], assumptions: Assumptions) -> None:
    k = result["kpis"]["monthly"]
    ws["A1"] = "NeuraFlow cash flow model"
    ws["A1"].font = Font(bold=True, size=16)
    rows = [
        ("Scenario", result["scenario_label"]),
        ("Forecast start", result["as_of"]),
        ("Forecast end", result["horizon_end"]),
        ("Starting cash", round(k["starting_cash"], 2)),
        ("Ending cash", round(k["ending_cash"], 2)),
        ("Net cash flow", round(k["net_cash_flow"], 2)),
        ("Operating cash flow", round(k["operating_cash_flow"], 2)),
        ("Investing cash flow", round(k["investing_cash_flow"], 2)),
        ("Financing cash flow", round(k["financing_cash_flow"], 2)),
        ("Lowest balance", round(k["lowest_balance"], 2)),
        ("Lowest balance date", k["lowest_balance_date"]),
        ("Minimum cash target", round(assumptions.general.min_cash, 2)),
        ("Funding gap vs minimum", round(k["funding_gap"], 2)),
        ("Runway (months)", k["runway_months"] if k["runway_months"] is not None else "Cash stays positive"),
    ]
    for r, (label, value) in enumerate(rows, start=3):
        ws.cell(row=r, column=1, value=label).font = BOLD
        cell = ws.cell(row=r, column=2, value=value)
        if isinstance(value, dt.date):
            cell.number_format = DATE
        elif isinstance(value, float):
            cell.number_format = MONEY
    r = 3 + len(rows) + 1
    ws.cell(row=r, column=1, value="Scenario adjustments applied").font = BOLD
    r += 1
    for key, value in result["effective_adjustments"].items():
        ws.cell(row=r, column=1, value=key.replace("_", " "))
        ws.cell(row=r, column=2, value=value)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Alerts").font = BOLD
    r += 1
    for alert in result["alerts"] or [{"message": "None"}]:
        ws.cell(row=r, column=1, value=alert["message"])
        r += 1
    r += 1
    ws.cell(
        row=r,
        column=1,
        value="Receivables come from the invoice data. Payroll, opex, debt and one-off items are only as good as the assumptions entered.",
    ).font = Font(italic=True)
    _widths(ws, 34, 24, 2)


def build_xlsx(result: Dict[str, Any], assumptions: Assumptions) -> bytes:
    wb = Workbook()
    _summary(wb.active, result, assumptions)
    wb.active.title = "Summary"
    g = assumptions.general
    _projections(wb.create_sheet("Projections"), result)
    _statement(wb.create_sheet("Monthly"), "Monthly cash flow statement", result["monthly"], g.starting_cash, g.min_cash)
    _statement(wb.create_sheet("13-Week"), "13-week cash flow", result["weekly"], g.starting_cash, g.min_cash)
    _pnl(wb.create_sheet("P&L"), result["pnl"])
    _scenarios(wb.create_sheet("Scenarios"), result)
    _ar(wb.create_sheet("Receivables"), result["ar"])
    _ap(wb.create_sheet("Payables"), result["ap"])
    _capex(wb.create_sheet("Capex"), result["capex"])
    if result.get("payroll"):
        _payroll(wb.create_sheet("Payroll"), result["payroll"])
    if result.get("balance_sheet"):
        _balance_sheet(wb.create_sheet("Balance Sheet"), result["balance_sheet"])
    if result.get("gl"):
        _gl(wb.create_sheet("GL vs Forecast"), result["gl"])
    _assumptions(wb.create_sheet("Assumptions"), assumptions)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_csv(result: Dict[str, Any]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    keys = [c[0] for c in CATEGORIES]
    labels = [c[1] for c in CATEGORIES]
    writer.writerow(
        ["period", "start", "end"] + labels
        + ["operating", "investing", "financing", "net_cash_flow", "beginning_cash", "ending_cash"]
    )
    for p in result["monthly"]:
        writer.writerow(
            [p["label"], p["start"].isoformat(), p["end"].isoformat()]
            + [round(p["categories"][k], 2) for k in keys]
            + [round(p[x], 2) for x in ("operating", "investing", "financing", "net", "begin_cash", "end_cash")]
        )
    return buffer.getvalue()
