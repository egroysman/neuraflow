"""Read tables from files, warehouses and SFTP. Nothing here writes to a source."""
from __future__ import annotations

import csv
import fnmatch
import gzip
import hashlib
import io
import os
import re
import zipfile
from dataclasses import dataclass
from typing import Any, Callable, Dict, List, Optional, Sequence, Tuple

ROW_LIMIT = int(os.getenv("INGEST_ROW_LIMIT", "250000"))
MAX_UPLOAD_BYTES = int(float(os.getenv("MAX_UPLOAD_MB", "25")) * 1024 * 1024)
MAX_UNZIPPED_BYTES = 300 * 1024 * 1024


class SourceError(Exception):
    """A problem the user can act on. The message is safe to show."""


Table = Tuple[List[str], List[Dict[str, Any]], bool]  # columns, rows, truncated


# ------------------------------------------------------------------ files

def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise SourceError("Could not read the file's text encoding.")


def clean_headers(headers: Sequence[Any]) -> List[str]:
    out: List[str] = []
    seen: Dict[str, int] = {}
    for i, h in enumerate(headers, 1):
        s = str(h).strip() if h is not None else ""
        s = s or f"Column {i}"
        if s in seen:
            seen[s] += 1
            s = f"{s} ({seen[s]})"
        else:
            seen[s] = 1
        out.append(s)
    return out


def detect_header_row(grid: Sequence[Sequence[Any]]) -> int:
    """0-based index of the header row: the first row that is mostly text with data under it."""
    for i, row in enumerate(grid[:15]):
        cells = [c for c in row if c is not None and str(c).strip() != ""]
        if len(cells) < 2:
            continue
        texty = sum(isinstance(c, str) and not re.fullmatch(r"[\d.,\-/$()\s%]+", c) for c in cells)
        if texty >= max(2, 0.6 * len(cells)) and any(any(c not in (None, "") for c in r) for r in grid[i + 1:i + 4]):
            return i
    return 0


def _grid_to_table(grid: List[List[Any]], header_row: Optional[int], limit: int) -> Table:
    if not grid:
        return [], [], False
    h = (header_row - 1) if header_row else detect_header_row(grid)
    cols = clean_headers(grid[h])
    rows: List[Dict[str, Any]] = []
    truncated = False
    for r in grid[h + 1:]:
        if not any(c is not None and str(c).strip() != "" for c in r):
            continue
        if len(rows) >= limit:
            truncated = True
            break
        rows.append({cols[j]: (r[j] if j < len(r) else None) for j in range(len(cols))})
    return cols, rows, truncated


def read_csv_bytes(raw: bytes, limit: int = ROW_LIMIT, header_row: Optional[int] = None) -> Table:
    if raw[:2] == b"\x1f\x8b":
        try:
            raw = gzip.decompress(raw)
        except OSError:
            raise SourceError("The file looks gzipped but could not be opened.")
    text = _decode(raw)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    grid = list(csv.reader(io.StringIO(text), dialect))
    return _grid_to_table(grid, header_row, limit)


def excel_sheets(raw: bytes) -> List[Dict[str, Any]]:
    wb = _open_workbook(raw)
    try:
        return [{"name": ws.title, "kind": "sheet", "detail": f"{ws.max_row or 0} rows" if ws.max_row else ""} for ws in wb.worksheets if ws.sheet_state == "visible"]
    finally:
        wb.close()


def _open_workbook(raw: bytes):
    if raw[:4] != b"PK\x03\x04":
        raise SourceError("This is not an .xlsx workbook. Old .xls files are not supported: open it in Excel and save as .xlsx or CSV.")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED_BYTES:
                raise SourceError("The workbook is too large when unpacked.")
    except zipfile.BadZipFile:
        raise SourceError("The workbook could not be opened. It may be damaged.")
    from openpyxl import load_workbook

    try:
        return load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise SourceError(f"The workbook could not be read: {type(e).__name__}.")


def read_excel_bytes(raw: bytes, sheet: Optional[str] = None, limit: int = ROW_LIMIT, header_row: Optional[int] = None) -> Table:
    wb = _open_workbook(raw)
    try:
        names = [w.title for w in wb.worksheets if w.sheet_state == "visible"]
        if not names:
            raise SourceError("The workbook has no visible sheets.")
        name = sheet if sheet in names else names[0]
        if sheet and sheet not in names:
            raise SourceError(f"Sheet '{sheet}' was not found. Sheets: {', '.join(names)}.")
        ws = wb[name]
        grid: List[List[Any]] = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            grid.append(list(row))
            if i > limit + 50:
                break
        return _grid_to_table(grid, header_row, limit)
    finally:
        wb.close()


def read_file_bytes(raw: bytes, filename: str, obj: Optional[str] = None, limit: int = ROW_LIMIT, header_row: Optional[int] = None) -> Table:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext in ("xlsx", "xlsm"):
        return read_excel_bytes(raw, obj, limit, header_row)
    if ext == "xls":
        raise SourceError("Old .xls files are not supported: open it in Excel and save as .xlsx or CSV.")
    if ext in ("csv", "tsv", "txt", "gz", ""):
        return read_csv_bytes(raw, limit, header_row)
    raise SourceError(f"Files of type .{ext} are not supported. Use CSV or Excel (.xlsx).")


def file_objects(raw: bytes, filename: str) -> List[Dict[str, Any]]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        return excel_sheets(raw)
    return [{"name": "file", "kind": "file", "detail": filename}]


# ------------------------------------------------------------------ warehouses

_IDENT = re.compile(r"^[A-Za-z0-9_$][A-Za-z0-9_$ \-]{0,127}$")
_FORBIDDEN_SQL = re.compile(r"\b(insert|update|delete|drop|alter|create|merge|grant|revoke|call|copy|put|get|truncate|replace|use|set|execute|exec|unset|remove|undrop|comment)\b", re.I)


def check_ident(part: str) -> str:
    if not _IDENT.match(part or ""):
        raise SourceError(f"'{part}' is not a valid table or schema name.")
    return part


def check_select(sql: str) -> str:
    s = (sql or "").strip().rstrip(";").strip()
    if not s:
        raise SourceError("Enter a query.")
    if ";" in s:
        raise SourceError("Use a single query with no semicolons.")
    if re.search(r"(--|/\*)", s):
        raise SourceError("Comments are not allowed in queries.")
    if not re.match(r"^(select|with)\b", s, re.I):
        raise SourceError("Only SELECT queries are allowed.")
    if _FORBIDDEN_SQL.search(s):
        raise SourceError("Only read-only queries are allowed.")
    return s


@dataclass
class Dialect:
    name: str
    quote_char: str

    def quote(self, part: str) -> str:
        return f"{self.quote_char}{check_ident(part)}{self.quote_char}"

    def table_ref(self, name: str, default_schema: Optional[str] = None, database: Optional[str] = None) -> str:
        parts = name.split(".")
        if len(parts) == 1 and default_schema:
            parts = [default_schema, parts[0]]
        if len(parts) > 3 or not parts:
            raise SourceError("Use SCHEMA.TABLE.")
        if len(parts) == 2 and database:
            parts = [database] + parts
        return ".".join(self.quote(p) for p in parts)

    def list_sql(self, database: Optional[str]) -> str:
        raise NotImplementedError


class SnowflakeDialect(Dialect):
    def __init__(self):
        super().__init__("snowflake", '"')

    def list_sql(self, database: Optional[str]) -> str:
        db = f"{self.quote(database)}." if database else ""
        return (f"SELECT table_schema, table_name, table_type, row_count FROM {db}information_schema.tables "
                "WHERE table_schema <> 'INFORMATION_SCHEMA' ORDER BY 1, 2")


class DatabricksDialect(Dialect):
    def __init__(self):
        super().__init__("databricks", "`")

    def list_sql(self, database: Optional[str]) -> str:
        cat = f"{self.quote(database)}." if database else ""
        return (f"SELECT table_schema, table_name, table_type, NULL AS row_count FROM {cat}information_schema.tables "
                "WHERE table_schema <> 'information_schema' ORDER BY 1, 2")


class DbReader:
    """A read-only reader over any DB-API connection factory."""

    def __init__(self, connect: Callable[[], Any], dialect: Dialect, database: Optional[str] = None, schema: Optional[str] = None):
        self.connect, self.dialect, self.database, self.schema = connect, dialect, database, schema

    def _run(self, sql: str, limit: Optional[int] = None):
        conn = self.connect()
        try:
            cur = conn.cursor()
            try:
                cur.execute(sql)
                cols = [str(d[0]) for d in (cur.description or [])]
                data = cur.fetchmany(limit + 1) if limit else cur.fetchall()
                return cols, data
            finally:
                try:
                    cur.close()
                except Exception:  # noqa: BLE001
                    pass
        finally:
            try:
                conn.close()
            except Exception:  # noqa: BLE001
                pass

    def test(self) -> None:
        try:
            self._run("SELECT 1")
        except SourceError:
            raise
        except Exception as e:  # noqa: BLE001
            raise SourceError(friendly_db_error(e))

    def objects(self) -> List[Dict[str, Any]]:
        try:
            _, data = self._run(self.dialect.list_sql(self.database))
        except SourceError:
            raise
        except Exception as e:  # noqa: BLE001
            raise SourceError(friendly_db_error(e))
        out = []
        for schema, name, ttype, count in data:
            if self.schema and str(schema).lower() != self.schema.lower():
                continue
            out.append({"name": f"{schema}.{name}", "kind": "view" if "VIEW" in str(ttype).upper() else "table",
                        "detail": f"{int(count):,} rows" if count not in (None, "") else ""})
        return out

    def read(self, obj: str, limit: int = ROW_LIMIT) -> Table:
        if obj.startswith("sql:"):
            inner = check_select(obj[4:])
            sql = f"SELECT * FROM ({inner}) AS q LIMIT {int(limit) + 1}"
        else:
            sql = f"SELECT * FROM {self.dialect.table_ref(obj, self.schema, self.database)} LIMIT {int(limit) + 1}"
        try:
            cols, data = self._run(sql)
        except SourceError:
            raise
        except Exception as e:  # noqa: BLE001
            raise SourceError(friendly_db_error(e))
        truncated = len(data) > limit
        cols = clean_headers(cols)
        return cols, [dict(zip(cols, r)) for r in data[:limit]], truncated


def friendly_db_error(e: Exception) -> str:
    msg = str(e).splitlines()[0][:240] if str(e) else type(e).__name__
    # never echo anything that looks like a secret
    msg = re.sub(r"(password|token|secret|key)\s*[=:]\s*\S+", r"\1=***", msg, flags=re.I)
    return f"The database said: {msg}"


def snowflake_reader(p: Dict[str, Any]) -> DbReader:
    for k in ("account", "user"):
        if not p.get(k):
            raise SourceError(f"Snowflake {k} is missing.")
    if not (p.get("password") or p.get("private_key")):
        raise SourceError("Give a password or a private key for Snowflake.")

    def connect():
        try:
            import snowflake.connector
        except ImportError:
            raise SourceError("The Snowflake driver is not installed on the server.")
        kw: Dict[str, Any] = dict(account=p["account"], user=p["user"], login_timeout=20, network_timeout=60)
        for k in ("warehouse", "database", "schema", "role"):
            if p.get(k):
                kw[k] = p[k]
        if p.get("private_key"):
            from cryptography.hazmat.primitives import serialization

            key = serialization.load_pem_private_key(p["private_key"].encode(), (p.get("private_key_passphrase") or "").encode() or None)
            kw["private_key"] = key.private_bytes(serialization.Encoding.DER, serialization.PrivateFormat.PKCS8, serialization.NoEncryption())
        else:
            kw["password"] = p["password"]
        kw["session_parameters"] = {"QUERY_TAG": "neuraflow-readonly"}
        return snowflake.connector.connect(**kw)

    return DbReader(connect, SnowflakeDialect(), p.get("database"), p.get("schema"))


def databricks_reader(p: Dict[str, Any]) -> DbReader:
    for k, label in (("server_hostname", "server hostname"), ("http_path", "HTTP path"), ("access_token", "access token")):
        if not p.get(k):
            raise SourceError(f"Databricks {label} is missing.")

    def connect():
        try:
            from databricks import sql as dbsql
        except ImportError:
            raise SourceError("The Databricks driver is not installed on the server.")
        kw: Dict[str, Any] = dict(server_hostname=p["server_hostname"], http_path=p["http_path"], access_token=p["access_token"])
        if p.get("catalog"):
            kw["catalog"] = p["catalog"]
        if p.get("schema"):
            kw["schema"] = p["schema"]
        return dbsql.connect(**kw)

    return DbReader(connect, DatabricksDialect(), p.get("catalog"), p.get("schema"))


# ------------------------------------------------------------------ SFTP

def host_key_fingerprint(key) -> str:
    import base64

    return "SHA256:" + base64.b64encode(hashlib.sha256(key.asbytes()).digest()).decode().rstrip("=")


class SftpReader:
    def __init__(self, p: Dict[str, Any], client_factory: Optional[Callable[[], Tuple[Any, Any]]] = None):
        for k in ("host", "username"):
            if not p.get(k):
                raise SourceError(f"SFTP {k} is missing.")
        if not (p.get("password") or p.get("private_key")):
            raise SourceError("Give a password or a private key for SFTP.")
        self.p = p
        self._factory = client_factory

    def _connect(self):
        """Returns (transport, sftp). The host key must match the pinned fingerprint unless told otherwise."""
        if self._factory:
            return self._factory()
        try:
            import paramiko
        except ImportError:
            raise SourceError("The SFTP driver is not installed on the server.")
        import socket

        try:
            sock = socket.create_connection((self.p["host"], int(self.p.get("port") or 22)), timeout=20)
            t = paramiko.Transport(sock)
            t.banner_timeout = 20
            t.start_client(timeout=20)
        except Exception as e:  # noqa: BLE001
            raise SourceError(f"Could not reach the SFTP server: {type(e).__name__}.")
        fp = host_key_fingerprint(t.get_remote_server_key())
        want = (self.p.get("host_key_sha256") or "").strip()
        if want and want != fp:
            t.close()
            raise SourceError("The server's host key does not match the one on file. Someone may be intercepting the connection, or the server was rebuilt.")
        if not want and not self.p.get("trust_any_host_key"):
            t.close()
            raise SourceError(f"HOSTKEY:{fp}")
        try:
            if self.p.get("private_key"):
                import io as _io

                pk = None
                for cls in (paramiko.Ed25519Key, paramiko.ECDSAKey, paramiko.RSAKey):
                    try:
                        pk = cls.from_private_key(_io.StringIO(self.p["private_key"]), password=self.p.get("private_key_passphrase") or None)
                        break
                    except Exception:  # noqa: BLE001
                        continue
                if pk is None:
                    raise SourceError("The SFTP private key could not be read.")
                t.auth_publickey(self.p["username"], pk)
            else:
                t.auth_password(self.p["username"], self.p["password"])
        except SourceError:
            t.close()
            raise
        except Exception as e:  # noqa: BLE001
            t.close()
            raise SourceError(f"SFTP sign-in failed: {type(e).__name__}.")
        return t, paramiko.SFTPClient.from_transport(t)

    def observed_fingerprint(self) -> Optional[str]:
        """Contact the server and return its host key fingerprint, without signing in."""
        import paramiko, socket  # noqa: E401

        try:
            sock = socket.create_connection((self.p["host"], int(self.p.get("port") or 22)), timeout=20)
            t = paramiko.Transport(sock)
            t.start_client(timeout=20)
            fp = host_key_fingerprint(t.get_remote_server_key())
            t.close()
            return fp
        except Exception as e:  # noqa: BLE001
            raise SourceError(f"Could not reach the SFTP server: {type(e).__name__}.")

    def test(self) -> None:
        t, c = self._connect()
        try:
            c.listdir(self.p.get("directory") or ".")
        except Exception as e:  # noqa: BLE001
            raise SourceError(f"Could not open the folder: {type(e).__name__}.")
        finally:
            c.close(); t.close()  # noqa: E702

    def objects(self) -> List[Dict[str, Any]]:
        t, c = self._connect()
        try:
            d = self.p.get("directory") or "."
            files = [a for a in c.listdir_attr(d) if not str(a.st_mode and oct(a.st_mode)).startswith("0o4")]
            files.sort(key=lambda a: a.st_mtime or 0, reverse=True)
            import datetime as dt

            return [{"name": a.filename, "kind": "file", "detail": f"{(a.st_size or 0) / 1024:,.0f} KB · " + dt.datetime.fromtimestamp(a.st_mtime or 0).strftime("%Y-%m-%d %H:%M")} for a in files if a.filename.lower().endswith((".csv", ".tsv", ".txt", ".gz", ".xlsx", ".xlsm"))]
        except SourceError:
            raise
        except Exception as e:  # noqa: BLE001
            raise SourceError(f"Could not list the folder: {type(e).__name__}.")
        finally:
            c.close(); t.close()  # noqa: E702

    def resolve(self, pattern: str) -> Tuple[str, bytes]:
        """Newest file in the folder matching a name or wildcard pattern. Returns (name, bytes)."""
        if "/" in pattern or "\\" in pattern or ".." in pattern:
            raise SourceError("Use a file name or a pattern like invoices*.csv, without folders.")
        t, c = self._connect()
        try:
            d = self.p.get("directory") or "."
            matches = [a for a in c.listdir_attr(d) if fnmatch.fnmatch(a.filename.lower(), pattern.lower())]
            if not matches:
                raise SourceError(f"No file matching '{pattern}' in {d}.")
            newest = max(matches, key=lambda a: a.st_mtime or 0)
            if (newest.st_size or 0) > MAX_UPLOAD_BYTES:
                raise SourceError(f"{newest.filename} is larger than the {MAX_UPLOAD_BYTES // (1024 * 1024)} MB limit.")
            path = (d.rstrip("/") + "/" + newest.filename) if d not in (".", "") else newest.filename
            with c.open(path, "rb") as fh:
                return newest.filename, fh.read(MAX_UPLOAD_BYTES + 1)
        except SourceError:
            raise
        except Exception as e:  # noqa: BLE001
            raise SourceError(f"Could not download the file: {type(e).__name__}.")
        finally:
            c.close(); t.close()  # noqa: E702

    def read(self, obj: str, limit: int = ROW_LIMIT, sheet: Optional[str] = None, header_row: Optional[int] = None) -> Table:
        name, raw = self.resolve(obj)
        if len(raw) > MAX_UPLOAD_BYTES:
            raise SourceError("The file is larger than the size limit.")
        return read_file_bytes(raw, name, sheet, limit, header_row)
