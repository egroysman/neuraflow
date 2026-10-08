"""Tests for payroll (roster + pay runs), the general ledger and the balance sheet."""
import datetime as dt
from collections import defaultdict

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from cashflow import ap, ar, engine, gl, payroll
from cashflow.models import (
    Adjustments, Assumptions, Capex, CapexItem, Employee, ForecastRequest, General, Loan, Macro,
    OpexLine, Payroll, Sales,
)
from cashflow.router import router

D = dt.date


def emp(id_="E1", hire="2025-01-06", term=None, salary=104000, dept="Ops", bonus=0, ben=0, **kw):
    return Employee(id=id_, department=dept, annual_salary=salary, hire_date=D.fromisoformat(hire),
                    term_date=D.fromisoformat(term) if term else None, bonus_pct=bonus, benefits_monthly=ben, **kw)


def assumptions(payroll_kwargs=None, **kw):
    return Assumptions(
        general=General(as_of=D(2026, 4, 9), horizon_months=12, starting_cash=200000, min_cash=0, tax_rate_pct=25),
        sales=Sales(monthly_revenue=300000, growth_pct_monthly=0.5, dso_days=40, bad_debt_pct=1),
        payroll=Payroll(**{"use_roster": True, "next_pay_date": D(2026, 4, 10), **(payroll_kwargs or {})}),
        opex=[OpexLine(name="Rent", kind="fixed", amount=20000)],
        loans=[Loan(name="Loan", balance=300000, annual_rate_pct=8, monthly_payment=8000)],
        **kw,
    )


def build_payroll(a):
    from cashflow.engine import add_months
    return payroll.build(a, 0.0, lambda k: (add_months(a.general.as_of, k), add_months(a.general.as_of, k + 1)))


# ----------------------------- sample data ---------------------------------


def test_sample_journal_balances_entry_by_entry():
    bal = defaultdict(float)
    for r in gl.load_gl()["journal"]:
        bal[r["EntryID"]] += float(r["Debit"]) - float(r["Credit"])
    assert all(abs(v) < 0.005 for v in bal.values())


def test_sample_gl_ties_to_invoices_and_bills():
    inv = ar.parse_invoices(ar.load_invoice_rows())
    bills = ap.parse_bills(ap.load_bill_rows())
    data = gl.overview(ar.snapshot_date(inv), inv, bills)
    assert data["trial_balance"]["balanced"]
    assert all(t["ok"] for t in data["tie_out"]), data["tie_out"]
    assert data["baselines"]["starting_cash"] > 0


def test_baselines_use_only_complete_months():
    p = gl.prepare(gl.load_gl())
    months = gl.monthly_actuals(p, D(2026, 4, 9))
    assert months[-1]["partial"] is True
    assert gl.baselines(p, D(2026, 4, 9))["window_months"] == ["2026-01", "2026-02", "2026-03"]


def test_gl_loan_baseline():
    loan = gl.baselines(gl.prepare(gl.load_gl()), D(2026, 4, 9))["loan"]
    assert loan["balance"] == pytest.approx(416000, abs=1)
    assert loan["monthly_payment"] == pytest.approx(9000, abs=1)
    assert loan["annual_rate_pct"] == pytest.approx(8, abs=0.05)


# ------------------------------- payroll -----------------------------------


def test_biweekly_has_26_or_27_runs_in_a_year():
    a = assumptions({"employees": [emp()]})
    runs = build_payroll(a)["runs"]
    assert len(runs) in (26, 27)
    assert runs[0]["date"] == D(2026, 4, 10)
    assert all((b["date"] - a["date"]).days == 14 for a, b in zip(runs, runs[1:]))


def test_semimonthly_and_monthly_run_counts():
    for freq, expected in (("semimonthly", 24), ("monthly", 12)):
        a = assumptions({"employees": [emp()], "pay_frequency": freq})
        assert abs(len(build_payroll(a)["runs"]) - expected) <= 1


def test_run_amount_matches_salary_over_periods():
    a = assumptions({"employees": [emp(salary=104000)], "employer_tax_pct": 10, "salary_growth_pct_annual": 0})
    run = build_payroll(a)["runs"][0]
    assert run["gross"] == pytest.approx(104000 / 26)
    assert run["total"] == pytest.approx(104000 / 26 * 1.10)


def test_hourly_pay():
    e = Employee(id="H", pay_type="hourly", hourly_rate=25, hours_per_week=40, hire_date=D(2025, 1, 1))
    assert e.annual_base() == 52000


def test_leaver_and_new_hire_are_prorated():
    a = assumptions({"employees": [emp("A", term="2026-06-30"), emp("B", hire="2026-08-17")], "salary_growth_pct_annual": 0})
    out = build_payroll(a)
    assert out["headcount"][0] == 1 and out["headcount"][-1] == 1
    assert max(out["headcount"]) == 1 or 2 in out["headcount"]
    assert out["monthly"][3] < out["monthly"][1] * 1.01  # July window: leaver gone, hire not yet
    assert out["monthly"][-1] > 0


def test_raises_step_in_raise_month():
    a = assumptions({"employees": [emp()], "raise_month": 1, "salary_growth_pct_annual": 5, "employer_tax_pct": 0})
    m = build_payroll(a)["monthly"]
    assert m[8] == pytest.approx(m[7] * 1.05, rel=0.02) or m[9] == pytest.approx(m[7] * 1.05, rel=0.02)
    assert m[3] == pytest.approx(m[0], rel=0.02)


def test_bonus_only_in_bonus_month_and_prorated_first_year():
    a = assumptions({"employees": [emp(bonus=10, hire="2025-01-06"), emp("N", bonus=10, hire="2026-01-05")], "bonus_month": 12,
                     "employer_tax_pct": 0, "salary_growth_pct_annual": 0})
    out = build_payroll(a)
    assert sum(1 for b in out["bonus"] if b > 0) == 1
    assert 10400 < sum(out["bonus"]) < 20800  # full year for one, partial for the new hire


def test_cash_and_accrual_totals_are_close():
    a = assumptions({"employees": [emp(bonus=5, ben=500), emp("B", salary=80000, ben=400)]})
    out = build_payroll(a)
    cash = sum(r["total"] for r in out["runs"]) + sum(out["bonus"]) * 1.085 + sum(out["benefits"])
    assert cash == pytest.approx(sum(out["monthly"]), rel=0.05)


def test_roster_mode_feeds_forecast_and_simple_mode_still_works():
    inv = ar.load_invoice_rows()
    roster = engine.forecast(inv, ForecastRequest(assumptions=assumptions({"employees": [emp()]})))
    assert roster["payroll"]["active_headcount"] == 1
    labels = [e.label for e in engine.run_model(ar.parse_invoices(inv), assumptions({"employees": [emp()]}), Adjustments()).events
              if e.category == "payroll"]
    assert labels and labels[0].startswith("Payroll run")
    simple = assumptions({"use_roster": False, "headcount": 5, "avg_salary": 80000})
    assert engine.forecast(inv, ForecastRequest(assumptions=simple))["payroll"] is None


def test_macro_inflation_increases_raises():
    base = assumptions({"employees": [emp()]})
    hot = assumptions({"employees": [emp()]}, macro=Macro(apply=True, cost_inflation_pct=5))
    inv = ar.load_invoice_rows()
    a = engine.forecast(inv, ForecastRequest(assumptions=base))["payroll"]["total_cost"]
    b = engine.forecast(inv, ForecastRequest(assumptions=hot))["payroll"]["total_cost"]
    assert b > a


# ----------------------------- balance sheet -------------------------------


def defaults_with_everything():
    rows = ar.load_invoice_rows()
    return rows, engine.default_assumptions(rows, bill_rows=ap.load_bill_rows(), payroll_rows=gl.load_payroll_rows(), gl_data=gl.load_gl())["assumptions"]


@pytest.mark.parametrize("scenario", ["base", "best", "worst"])
def test_balance_sheet_balances_in_every_scenario(scenario):
    rows, a = defaults_with_everything()
    r = engine.forecast(rows, ForecastRequest(assumptions=a, scenario=scenario), ap.load_bill_rows(), gl.load_gl())
    assert r["balance_sheet"]["max_abs_check"] < 0.01
    assert len(r["balance_sheet"]["months"]) == a.general.horizon_months


def test_balance_sheet_balances_with_adjustments_macro_and_financed_capex():
    rows, a = defaults_with_everything()
    a.macro = Macro(apply=True, rate_change_pts=1.5, cost_inflation_pct=4, demand_growth_pct=-3)
    a.capex.items.append(CapexItem(name="Truck", category="vehicles", date=D(2026, 7, 1), amount=90000, funding="lease", term_months=24))
    adj = Adjustments(collection_delay_days=20, extra_bad_debt_pct=3, revenue_change_pct=-15, capex_change_pct=50, dpo_change_days=-10, opex_change_pct=10)
    r = engine.forecast(rows, ForecastRequest(assumptions=a, scenario="worst", adjustments=adj), ap.load_bill_rows(), gl.load_gl())
    assert r["balance_sheet"]["max_abs_check"] < 0.01


def test_balance_sheet_with_one_time_items_and_no_bills():
    from cashflow.models import OneTimeItem
    a = assumptions({"employees": [emp()]})
    a.one_time = [OneTimeItem(name="Sale of van", date=D(2026, 6, 1), amount=15000, category="investing"),
                  OneTimeItem(name="Refund", date=D(2026, 7, 1), amount=-4000, category="operating"),
                  OneTimeItem(name="Owner draw", date=D(2026, 8, 1), amount=-10000, category="financing")]
    a.capex = Capex(opening_ppe_net=100000, existing_depreciation_monthly=2000)
    a.costs.opening_ap = 30000
    r = engine.forecast(ar.load_invoice_rows(), ForecastRequest(assumptions=a))
    assert r["balance_sheet"]["max_abs_check"] < 0.01
    assert r["balance_sheet"]["opening"]["payables"] == 30000


def test_balance_sheet_opening_and_memo():
    rows, a = defaults_with_everything()
    bs = engine.forecast(rows, ForecastRequest(assumptions=a), ap.load_bill_rows(), gl.load_gl())["balance_sheet"]
    assert bs["opening"]["receivables"] == pytest.approx(1056347.10, abs=0.5)
    assert bs["opening"]["payables"] == pytest.approx(222182.55, abs=0.5)
    assert bs["opening"]["cash"] == a.general.starting_cash
    assert bs["memo"]["existing_ar_expected_uncollectible"] > 0


# --------------------------------- GL drive --------------------------------


def test_defaults_driven_by_gl():
    rows = ar.load_invoice_rows()
    d = engine.default_assumptions(rows, bill_rows=ap.load_bill_rows(), payroll_rows=gl.load_payroll_rows(), gl_data=gl.load_gl())
    a, summary = d["assumptions"], d["data_summary"]
    assert {"starting_cash", "monthly_revenue", "cogs_pct", "opex", "loans"} <= set(summary["driven_by_gl"])
    assert a.general.starting_cash == pytest.approx(520000, abs=5)
    assert a.payroll.use_roster and len(a.payroll.employees) == 16
    assert a.loans[0].balance == 416000


def test_defaults_without_gl_are_unchanged_placeholders():
    d = engine.default_assumptions(ar.load_invoice_rows(), bill_rows=ap.load_bill_rows())
    assert d["data_summary"]["driven_by_gl"] == []
    assert not d["assumptions"].payroll.use_roster
    assert d["assumptions"].capex.opening_ppe_net == 0


def test_forecast_compare_has_timeline_and_baseline_check():
    rows, a = defaults_with_everything()
    r = engine.forecast(rows, ForecastRequest(assumptions=a), ap.load_bill_rows(), gl.load_gl())["gl"]
    tl = r["timeline"]
    assert sum(t["forecast_revenue"] is not None for t in tl) == 12
    assert 1 <= sum(t["actual_revenue"] is not None for t in tl) <= 12
    assert all(t["actual_revenue"] is None for t in tl[-12:])  # actuals end before the forecast starts
    assert r["baseline_check"][0]["label"] == "Revenue"
    assert abs(r["baseline_check"][0]["difference_pct"]) < 1.0  # defaults come from the ledger


def test_variance_appears_when_start_date_is_moved_back():
    rows, a = defaults_with_everything()
    a.general.as_of = D(2026, 1, 1)
    r = engine.forecast(rows, ForecastRequest(assumptions=a), ap.load_bill_rows(), gl.load_gl())["gl"]
    assert r["variance"], "forecast months that already have actuals should be compared"
    assert r["variance"][0]["lines"][0]["label"] == "Revenue"


def test_forecast_without_gl_returns_none():
    rows, a = defaults_with_everything()
    assert engine.forecast(rows, ForecastRequest(assumptions=a), ap.load_bill_rows())["gl"] is None


# --------------------------------- API -------------------------------------


@pytest.fixture
def client():
    app = FastAPI()
    app.include_router(router)
    return TestClient(app)


def test_api_gl_endpoint(client):
    body = client.get("/cashflow/gl").json()
    assert body["trial_balance"]["balanced"]
    assert all(t["ok"] for t in body["tie_out"])
    assert body["accounts"] and body["monthly"] and body["recent_entries"]


def test_api_defaults_and_forecast_include_new_sections(client):
    d = client.get("/cashflow/defaults").json()
    assert d["data_summary"]["driven_by_gl"]
    body = client.post("/cashflow/forecast", json={"assumptions": d["assumptions"], "scenario": "base", "adjustments": {}}).json()
    assert body["balance_sheet"]["max_abs_check"] < 0.01
    assert body["payroll"]["employees"] and body["gl"]["timeline"]


def test_xlsx_export_contains_new_sheets_and_balance_check_row(client):
    import io
    from openpyxl import load_workbook
    a = client.get("/cashflow/defaults").json()["assumptions"]
    wb = load_workbook(io.BytesIO(client.post("/cashflow/export?format=xlsx", json={"assumptions": a}).content))
    assert {"Payroll", "Balance Sheet", "GL vs Forecast"} <= set(wb.sheetnames)
    ws = wb["Balance Sheet"]
    check = [c.value for c in ws[15][1:]]  # the check row
    assert ws.cell(row=15, column=1).value.startswith("Check") and all(abs(v) < 0.02 for v in check)


def test_accrued_payroll_never_negative_and_opening_liability():
    rows, a = defaults_with_everything()
    r = engine.forecast(rows, ForecastRequest(assumptions=a), ap.load_bill_rows(), gl.load_gl())
    bs = r["balance_sheet"]
    assert bs["opening"]["accrued_payroll"] > 0
    assert bs["max_abs_check"] < 0.01
    assert min(m["accrued_payroll"] for m in bs["months"]) > -0.01 * r["payroll"]["total_cost"] / 12 * 5
